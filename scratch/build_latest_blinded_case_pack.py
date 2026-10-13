import json
import re
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Automation\Adjudication app")
SOURCE = ROOT / "scratch" / "latest_blinded_presentations.json"
OUT = ROOT / "output" / "docx" / "ACRN_Latest_Blinded_Cases_Presentation_and_Laboratory_Review_Pack.docx"
OUT.parent.mkdir(parents=True, exist_ok=True)

NAVY = "17365D"
PALE = "F4F7FB"
PALE_BLUE = "EAF1F8"
GRAY = "D9D9D9"
INK = "1F2937"
MUTED = "5B6573"
WHITE = "FFFFFF"
AMBER = "FFF2CC"

LAB_UNITS = {
    "HAEMOGLOBIN": "g/dL",
    "HAEMATOCRIT": "%",
    "WBC": "x10^9/L",
    "RBC": "x10^12/L",
    "PLATELETS": "x10^9/L",
    "ANC": "x10^9/L",
    "ALC": "x10^9/L",
    "CREATININE": "umol/L",
    "AST": "U/L",
    "ALT": "U/L",
    "ALP": "U/L",
    "LDH": "U/L",
    "TOTAL_BILIRUBIN": "umol/L",
    "BUN": "mmol/L",
}


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=GRAY, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_cell_margins(cell, top=70, start=85, bottom=70, end=85):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def style_run(run, size=9, bold=False, color=INK, font="Aptos"):
    run.font.name = font
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:ascii"), font)
    rpr.rFonts.set(qn("w:hAnsi"), font)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(9.2)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    title = doc.styles["Title"]
    title.font.name = "Aptos Display"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    title.font.size = Pt(26)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(9)
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for name, size, before, after in (("Heading 1", 17, 9, 6), ("Heading 2", 12.5, 8, 4), ("Heading 3", 10.2, 6, 3)):
        style = doc.styles[name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_table(doc, headers, rows, widths=None, font_size=8.0, header_fill=NAVY, compact=False):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header = table.rows[0]
    repeat_header(header)
    for index, (cell, label) in enumerate(zip(header.cells, headers)):
        set_cell_shading(cell, header_fill)
        set_cell_border(cell)
        set_cell_margins(cell, 65 if compact else 85, 70 if compact else 95, 65 if compact else 85, 70 if compact else 95)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        style_run(paragraph.add_run(str(label)), size=font_size, bold=True, color=WHITE)
        if widths:
            cell.width = Inches(widths[index])
    for row_index, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for index, (cell, value) in enumerate(zip(row.cells, values)):
            if row_index % 2:
                set_cell_shading(cell, PALE)
            set_cell_border(cell)
            set_cell_margins(cell, 55 if compact else 75, 65 if compact else 85, 55 if compact else 75, 65 if compact else 85)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            style_run(paragraph.add_run(str(value)), size=font_size)
            if widths:
                cell.width = Inches(widths[index])
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(1)
    return table


def add_labelled_paragraph(doc, label, text, size=9.2):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(5)
    style_run(paragraph.add_run(label + " "), size=size, bold=True)
    style_run(paragraph.add_run(str(text)), size=size)
    return paragraph


def add_narrative(doc, text, size=9.1):
    blocks = [block.strip() for block in str(text or "").split("\n\n") if block.strip()]
    if not blocks:
        blocks = ["No presentation was generated."]
    for block in blocks:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(6)
        paragraph.paragraph_format.line_spacing = 1.12
        style_run(paragraph.add_run(block), size=size)


def add_bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.left_indent = Inches(0.22)
        paragraph.paragraph_format.first_line_indent = Inches(-0.12)
        paragraph.paragraph_format.space_after = Pt(2)
        style_run(paragraph.add_run(item), size=8.8)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    style_run(run, size=8, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def fmt_date(value, with_time=False):
    if not value:
        return "Unavailable"
    try:
        dt = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
        return dt.strftime("%d %b %Y %H:%M" if with_time else "%d %b %Y")
    except ValueError:
        return str(value)


def unique(values):
    result = []
    for value in values:
        text = str(value).strip()
        if text and text not in result:
            result.append(text)
    return result


def visit_bp(visit):
    values = []
    for item in visit.get("bp", []):
        if item.get("sbp") is None or item.get("dbp") is None:
            continue
        label = f"{item['sbp']}/{item['dbp']}"
        if item.get("kind") == "recheck":
            label += " recheck"
        values.append(label)
    return "; ".join(unique(values)) or "Not documented"


def visit_protein(visit):
    values = []
    for item in visit.get("proteinuria", []):
        value = item.get("value")
        if value in (None, ""):
            continue
        unit = f" {item.get('unit')}" if item.get("unit") else ""
        values.append(f"{item.get('method') or 'Protein'}: {value}{unit}")
    return "; ".join(unique(values)) or "Not documented"


def source_row(label):
    match = re.search(r"/ row (\d+)$", str(label or ""))
    return f"r{match.group(1)}" if match else ""


def lab_cell(visit, key):
    results = []
    for lab in visit.get("labs", []):
        if lab.get("key") != key:
            continue
        display = str(lab.get("display") or lab.get("raw") or lab.get("value") or "Not documented")
        interpretation = (lab.get("interpretation") or {}).get("label")
        row = source_row(lab.get("source_label"))
        suffix = []
        if interpretation and interpretation not in {"Normal", "Not documented"}:
            suffix.append(interpretation)
        if row:
            suffix.append(row)
        if suffix:
            display += " [" + ", ".join(suffix) + "]"
        results.append(display)
    return "; ".join(unique(results)) or "-"


def case_gp(case):
    risk = case.get("risk_summary") or {}
    return f"G{risk.get('gravidity', 0)}P{risk.get('parity', 0)}"


def case_index_row(case):
    visits = case.get("normalized_visits", [])
    available_labs = sum(len(visit.get("labs", [])) for visit in visits)
    return [
        case.get("subject_id"),
        case.get("age") if case.get("age") is not None else "Unavailable",
        case_gp(case),
        len(visits),
        f"{fmt_date(case.get('first_visit'))} to {fmt_date(case.get('latest_visit'))}",
        available_labs,
        case.get("qc_status") or "Unavailable",
    ]


payload = json.loads(SOURCE.read_text(encoding="utf-8"))
batch = payload["batch"]
cases = payload["cases"]

doc = Document()
section = doc.sections[0]
section.orientation = WD_ORIENT.LANDSCAPE
section.page_width, section.page_height = section.page_height, section.page_width
section.top_margin = Inches(0.52)
section.bottom_margin = Inches(0.52)
section.left_margin = Inches(0.58)
section.right_margin = Inches(0.58)
configure_styles(doc)

footer_table = section.footer.add_table(rows=1, cols=2, width=Inches(10.4))
footer_table.columns[0].width = Inches(8.6)
footer_table.columns[1].width = Inches(1.8)
footer_text = footer_table.cell(0, 0).paragraphs[0]
style_run(footer_text.add_run("ACRN Blinded Case Presentation and Laboratory Review Pack  |  UAT working document"), size=8, color=MUTED)
add_page_number(footer_table.cell(0, 1).paragraphs[0])

title = doc.add_paragraph(style="Title")
title.add_run("Blinded Case Presentation and Laboratory Review Pack")
subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(13)
style_run(subtitle.add_run("Latest Monitor upload and Adjudicator presentation output"), size=14, color=MUTED)

intro = doc.add_paragraph()
intro.paragraph_format.space_after = Pt(8)
style_run(
    intro.add_run(
        "This editable review pack is generated from the latest blinded RealTime batch currently loaded in the application. "
        "It reproduces the application-aligned overall and per-visit presentations and places the displayed laboratory results beside the source visit chronology. "
        "Reviewers can redline mapping, chronology, missing-data language and presentation wording before the rerun is linked to the existing UAT record."
    ),
    size=10.1,
)

add_table(doc, ["Document field", "Controlled value"], [
    ["Source upload", batch.get("filename")],
    ["Upload timestamp", fmt_date(batch.get("uploaded_at"), with_time=True)],
    ["Blinded participants", len(cases)],
    ["Source rows", batch.get("row_count")],
    ["Reconstructed source visits", batch.get("visit_count")],
    ["Mapping version recorded at upload", batch.get("mapping_version")],
    ["Narrative version", payload.get("narrative_version")],
    ["Import status", batch.get("status")],
    ["Use", "Clinical presentation correction, source mapping review and UAT rerun evidence"],
], widths=[2.45, 7.9], font_size=8.8)

doc.add_heading("How to review this pack", level=1)
add_bullets(doc, [
    "Confirm that each visit date and gestational age belongs to that visit and follows the approved dating anchor.",
    "Confirm blood-pressure pairing, repeat readings, proteinuria method and missing-data wording against the displayed source evidence.",
    "Redline the overall and per-visit presentations where the wording is clinically inaccurate, incomplete or insufficiently neutral.",
    "Check every laboratory cell against its source row marker. A source marker such as r2449 refers to that row in the uploaded source file.",
    "Record the corrected expected wording and mapping rule before the application rerun; screenshots remain supporting evidence only.",
])

doc.add_heading("Application presentation model", level=1)
add_table(doc, ["Displayed element", "Application source", "Review requirement"], [
    ["Participant header", "Blinded subject identifier, age and derived obstetric history", "No direct identifiers; age and G/P agree with source history"],
    ["Visit chronology", "Reconstructed V01 to V06 visit instances", "Dates and GA remain visit-specific; postpartum is not represented as antenatal GA"],
    ["Blood pressure", "Paired systolic and diastolic evidence including repeats", "No cross-row or cross-visit pairing"],
    ["Proteinuria", "Dipstick and quantitative protein evidence", "Method and actual result remain visible"],
    ["Laboratory display", "Urinalysis, haematology and biochemistry result evidence", "Value, interpretation and source-row marker remain traceable"],
    ["Visit presentation", "FORM-ADJ-PROSE v2.0 presentation generator", "Neutral clinical prose; unavailable evidence is explicit"],
    ["Overall presentation", "Chronological synthesis across V01 to V06", "No automated final adjudication; unresolved source issues remain visible"],
], widths=[1.7, 4.0, 4.65], font_size=8.3)

doc.add_page_break()
doc.add_heading("Blinded case index", level=1)
add_labelled_paragraph(doc, "Batch scope", f"All {len(cases)} blinded participants from {batch.get('filename')} are included in source order.")
add_table(doc, ["Blinded subject", "Age", "G/P", "V01-V06", "Source date range", "Displayed labs", "Workflow status"], [case_index_row(case) for case in cases], widths=[1.65, 0.55, 0.65, 0.7, 2.3, 1.0, 1.75], font_size=7.7)

doc.add_heading("Laboratory conventions", level=1)
add_labelled_paragraph(doc, "Display", "Each result is shown exactly as normalized for the application. A non-normal interpretation is shown in square brackets, followed by the uploaded source-row marker.")
add_labelled_paragraph(doc, "Missingness", "A dash means that the application has no displayed result for that analyte at that visit. It must not be interpreted as normal, negative or zero.")
add_labelled_paragraph(doc, "Units", "Units follow the application laboratory metadata. Reviewers should confirm them against the study laboratory specification before approval.")

for case_number, case in enumerate(cases, start=1):
    visits = case.get("normalized_visits", [])
    risk = case.get("risk_summary") or {}
    doc.add_page_break()
    doc.add_heading(f"Case {case_number:02d}  {case.get('subject_id')}", level=1)
    add_table(doc, ["Age", "Obstetric history", "Source visits", "Source date range", "Import readiness", "Open issues"], [[
        case.get("age") if case.get("age") is not None else "Unavailable",
        case_gp(case),
        len(visits),
        f"{fmt_date(case.get('first_visit'))} to {fmt_date(case.get('latest_visit'))}",
        (case.get("import_readiness") or {}).get("status") or "Unavailable",
        case.get("open_issues", 0),
    ]], widths=[0.65, 1.2, 0.9, 2.5, 1.7, 0.8], font_size=8.0)
    add_labelled_paragraph(doc, "History summary", f"Gravidity {risk.get('gravidity', 0)}; parity {risk.get('parity', 0)}; miscarriages {risk.get('miscarriages', 0)}; stillbirths {risk.get('stillbirths', 0)}; chronic hypertension {'recorded' if risk.get('chronic_htn') else 'not recorded'}.")

    doc.add_heading("Overall presentation from the Adjudicator view", level=2)
    add_narrative(doc, case.get("overall_presentation"), size=9.0)

    doc.add_heading("Source chronology", level=2)
    chronology_rows = []
    for visit in visits:
        chronology_rows.append([
            visit.get("name") or f"V{visit.get('visit_number', 0):02d}",
            fmt_date(visit.get("date")),
            visit.get("ga") or ("Postpartum" if visit.get("ga_provenance") == "POSTPARTUM" else "Unavailable"),
            visit_bp(visit),
            visit_protein(visit),
            len(visit.get("labs", [])),
            (visit.get("interpretation") or {}).get("classification") or "Not assessable",
        ])
    add_table(doc, ["Visit", "Date", "Gestational age", "Blood pressure", "Proteinuria", "Labs", "Application interpretation"], chronology_rows, widths=[0.6, 1.05, 1.5, 1.65, 2.25, 0.55, 2.7], font_size=7.5, compact=True)

    doc.add_page_break()
    doc.add_heading(f"{case.get('subject_id')}  Per visit presentations", level=1)
    presentations = {int(item.get("visit_number") or 0): item.get("presentation") for item in case.get("visit_presentations", [])}
    for visit in visits:
        number = int(visit.get("visit_number") or 0)
        label = visit.get("name") or f"V{number:02d}"
        doc.add_heading(f"{label}  {fmt_date(visit.get('date'))}  {visit.get('ga') or 'GA unavailable'}", level=2)
        add_narrative(doc, presentations.get(number), size=8.8)

    doc.add_page_break()
    doc.add_heading(f"{case.get('subject_id')}  Laboratory results", level=1)
    add_labelled_paragraph(doc, "Cell format", "Result [non-normal interpretation, source row]. A dash means no displayed result for that visit.", size=8.7)
    visit_headers = [visit.get("name") or f"V{int(visit.get('visit_number') or 0):02d}" for visit in visits]
    visit_width = 1.42 if len(visits) == 6 else max(0.9, 8.5 / max(1, len(visits)))
    for category in ("Urinalysis", "Haematology", "Biochemistry"):
        category_labs = []
        seen_keys = set()
        for visit in visits:
            for lab in visit.get("labs", []):
                if lab.get("category") == category and lab.get("key") not in seen_keys:
                    category_labs.append((lab.get("key"), lab.get("label") or lab.get("key")))
                    seen_keys.add(lab.get("key"))
        if not category_labs:
            continue
        doc.add_heading(category, level=2)
        rows = []
        for key, label in category_labs:
            unit = LAB_UNITS.get(key)
            analyte = f"{label} ({unit})" if unit else label
            rows.append([analyte] + [lab_cell(visit, key) for visit in visits])
        add_table(doc, ["Analyte"] + visit_headers, rows, widths=[1.85] + [visit_width] * len(visits), font_size=6.8, compact=True)

    doc.add_heading("Presentation and mapping corrections", level=2)
    add_table(doc, ["Review item", "Expected correction or approved wording", "Owner and disposition"], [
        ["Overall presentation", "", ""],
        ["Visit chronology and gestational age", "", ""],
        ["Blood pressure and proteinuria mapping", "", ""],
        ["Laboratory values, units and interpretation", "", ""],
        ["Missing or conflicting source evidence", "", ""],
    ], widths=[2.6, 5.25, 2.35], font_size=8.0)

doc.add_page_break()
doc.add_heading("Rerun evidence and sign off record", level=1)
add_labelled_paragraph(doc, "Instruction", "Complete one row for every blinded case after corrections are implemented and rerun. Link the full output and source mapping evidence into the existing UAT record with Panashe.")
add_table(doc, ["Blinded subject", "Mapping result", "Presentation result", "Laboratory result", "Open defect", "Reviewer and date"], [[case.get("subject_id"), "Not run", "Not run", "Not run", "", ""] for case in cases], widths=[1.65, 1.7, 1.8, 1.65, 1.9, 1.55], font_size=7.6)

doc.add_heading("Required rerun evidence", level=1)
add_bullets(doc, [
    "Editable reviewed template with Tariro's accepted corrections.",
    "Field-to-source mapping output that identifies the source field, source row, canonical field and transformation rule.",
    "Complete overall and per-visit presentation outputs for every blinded case in this batch.",
    "Complete laboratory outputs with values, units, interpretation, visit date and source reference.",
    "Expected-versus-actual results and point-by-point responses for every open defect.",
    "Recorded clinical review and implementation-clearance decision linked to the existing UAT record with Panashe.",
])

doc.core_properties.title = "Blinded Case Presentation and Laboratory Review Pack"
doc.core_properties.subject = "Latest application batch presentations laboratories source chronology and UAT review"
doc.core_properties.author = "ACRN Adjudication Programme"
doc.core_properties.keywords = "blinded cases, adjudication, monitor mapping, presentations, laboratory results, UAT"
doc.save(OUT)
print(OUT)
